"""Automated tests for PDF, DOCX, and TXT extraction."""

import tempfile
import unittest
from pathlib import Path

import pymupdf as fitz
from docx import Document

from src.document_intelligence.ingestion import (
    DocumentExtractionError,
    UnsupportedDocumentError,
    extract_document,
)


class ExtractDocumentTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.directory = Path(self.temp_dir.name)

    def test_extracts_txt_and_preserves_filename(self) -> None:
        path = self.directory / "notes.txt"
        path.write_text("First line.\nSecond line.", encoding="utf-8")

        result = extract_document(path)

        self.assertEqual(result.source_filename, "notes.txt")
        self.assertEqual(result.text, "First line.\nSecond line.")
        self.assertEqual(len(result.pages), 1)
        self.assertIsNone(result.pages[0].page_number)
        self.assertFalse(result.is_empty)

    def test_extracts_docx_paragraphs_and_table_without_page_number(self) -> None:
        path = self.directory / "report.docx"
        document = Document()
        document.add_paragraph("Report introduction.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Value"
        document.add_paragraph("Report conclusion.")
        document.save(path)

        result = extract_document(path)

        self.assertEqual(result.source_filename, "report.docx")
        self.assertEqual(
            result.text,
            "Report introduction.\nName\tValue\nReport conclusion.",
        )
        self.assertEqual(len(result.pages), 1)
        self.assertIsNone(result.pages[0].page_number)

    def test_extracts_pdf_text_with_one_based_page_numbers(self) -> None:
        path = self.directory / "manual.pdf"
        pdf = fitz.open()
        first_page = pdf.new_page()
        first_page.insert_text((72, 72), "First page text.")
        pdf.new_page()
        third_page = pdf.new_page()
        third_page.insert_text((72, 72), "Third page text.")
        pdf.save(path)
        pdf.close()

        result = extract_document(path)

        self.assertEqual(result.source_filename, "manual.pdf")
        self.assertEqual([page.page_number for page in result.pages], [1, 2, 3])
        self.assertIn("First page text.", result.pages[0].text)
        self.assertEqual(result.pages[1].text.strip(), "")
        self.assertIn("Third page text.", result.pages[2].text)
        self.assertFalse(result.is_empty)

    def test_empty_and_whitespace_only_documents_are_marked_empty(self) -> None:
        txt_path = self.directory / "empty.txt"
        txt_path.write_text(" \t\r\n ", encoding="utf-8")

        docx_path = self.directory / "empty.docx"
        Document().save(docx_path)

        pdf_path = self.directory / "empty.pdf"
        pdf = fitz.open()
        pdf.new_page()
        pdf.save(pdf_path)
        pdf.close()

        for path in (txt_path, docx_path, pdf_path):
            with self.subTest(extension=path.suffix):
                result = extract_document(path)
                self.assertTrue(result.is_empty)

    def test_unsupported_extension_has_helpful_error(self) -> None:
        path = self.directory / "archive.csv"
        path.write_text("a,b", encoding="utf-8")

        with self.assertRaisesRegex(
            UnsupportedDocumentError, "Supported extensions: .pdf, .docx, .txt"
        ):
            extract_document(path)

    def test_missing_file_has_helpful_error(self) -> None:
        with self.assertRaisesRegex(DocumentExtractionError, "not an existing"):
            extract_document(self.directory / "missing.txt")

    def test_corrupt_pdf_and_docx_are_not_reported_as_success(self) -> None:
        pdf_path = self.directory / "broken.pdf"
        pdf_path.write_bytes(b"not a PDF")
        docx_path = self.directory / "broken.docx"
        docx_path.write_bytes(b"not a DOCX")

        for path in (pdf_path, docx_path):
            with self.subTest(extension=path.suffix):
                with self.assertRaisesRegex(
                    DocumentExtractionError, "Could not extract"
                ):
                    extract_document(path)

    def test_invalid_utf8_text_has_helpful_error(self) -> None:
        path = self.directory / "invalid.txt"
        path.write_bytes(b"\xff\xfe\xfa")

        with self.assertRaisesRegex(DocumentExtractionError, "as UTF-8"):
            extract_document(path)


if __name__ == "__main__":
    unittest.main()
