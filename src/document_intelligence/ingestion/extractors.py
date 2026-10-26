"""Extract text and source metadata from PDF, DOCX, and TXT documents."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

import pymupdf as fitz
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph


@dataclass(frozen=True)
class ExtractedPage:
    """Text from one source page, when the file format has page boundaries.

    ``page_number`` is one-based for PDF pages. It is ``None`` for DOCX and TXT
    files because those formats do not provide reliable page boundaries.
    """

    page_number: int | None
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    """Extracted text with the original filename and its source sections.

    ``pages`` contains one item per PDF page, including blank pages. DOCX and
    TXT documents contain one item with ``page_number=None``. ``text`` joins
    those items in source order with newline separators. ``is_empty`` is true
    when the resulting text is empty or whitespace-only.
    """

    source_filename: str
    pages: tuple[ExtractedPage, ...]

    @property
    def text(self) -> str:
        """Return all extracted text in source order."""
        return "\n".join(page.text for page in self.pages)

    @property
    def is_empty(self) -> bool:
        """Indicate whether extraction produced only empty or whitespace text."""
        return not self.text.strip()


class DocumentExtractionError(Exception):
    """Raised when a supported document cannot be read or extracted."""


class UnsupportedDocumentError(DocumentExtractionError):
    """Raised when a file extension is not supported."""


def extract_document(file_path: str | Path) -> ExtractedDocument:
    """Extract a PDF, DOCX, or UTF-8 TXT file.

    Empty and whitespace-only documents are returned with ``is_empty=True``.
    Unsupported formats and files that cannot be read raise a
    ``DocumentExtractionError`` subclass with a user-facing explanation.
    """
    path = Path(file_path)
    extension = path.suffix.lower()
    extractors = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".txt": _extract_txt,
    }

    try:
        extractor = extractors[extension]
    except KeyError as exc:
        supported = ", ".join(extractors)
        raise UnsupportedDocumentError(
            f"Unsupported file extension '{extension or '(none)'}' for "
            f"'{path.name}'. Supported extensions: {supported}."
        ) from exc

    if not path.is_file():
        raise DocumentExtractionError(
            f"'{path}' is not an existing regular file."
        )

    return extractor(path)


def extract_document_content(
    source_filename: str,
    content: bytes,
) -> ExtractedDocument:
    """Extract supported document bytes without writing an uploaded file.

    This entry point is suitable for in-memory uploads. The filename is used
    for format detection and retained as the document's source metadata.
    """
    if not isinstance(content, bytes):
        raise TypeError("document content must be bytes")

    path = Path(source_filename)
    extension = path.suffix.lower()
    if extension == ".pdf":
        return _extract_pdf_bytes(path.name, content)
    if extension == ".docx":
        return _extract_docx_bytes(path.name, content)
    if extension == ".txt":
        return _extract_txt_bytes(path.name, content)

    supported = ".pdf, .docx, .txt"
    raise UnsupportedDocumentError(
        f"Unsupported file extension '{extension or '(none)'}' for "
        f"'{path.name}'. Supported extensions: {supported}."
    )


def _extract_pdf(path: Path) -> ExtractedDocument:
    try:
        content = path.read_bytes()
    except OSError as exc:
        raise DocumentExtractionError(
            f"Could not extract PDF '{path.name}': {exc}"
        ) from exc
    return _extract_pdf_bytes(path.name, content)


def _extract_pdf_bytes(filename: str, content: bytes) -> ExtractedDocument:
    try:
        with fitz.open(stream=content, filetype="pdf") as pdf:
            pages = tuple(
                ExtractedPage(page_number=index, text=page.get_text())
                for index, page in enumerate(pdf, start=1)
            )
    except (fitz.FileDataError, OSError, RuntimeError, ValueError) as exc:
        raise DocumentExtractionError(
            f"Could not extract PDF '{filename}': {exc}"
        ) from exc

    return ExtractedDocument(source_filename=filename, pages=pages)


def _extract_docx(path: Path) -> ExtractedDocument:
    try:
        content = path.read_bytes()
    except OSError as exc:
        raise DocumentExtractionError(
            f"Could not extract DOCX '{path.name}': {exc}"
        ) from exc
    return _extract_docx_bytes(path.name, content)


def _extract_docx_bytes(filename: str, content: bytes) -> ExtractedDocument:
    try:
        document = Document(BytesIO(content))
        paragraphs_and_rows: list[str] = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                paragraphs_and_rows.append(block.text)
            elif isinstance(block, Table):
                paragraphs_and_rows.extend(
                    "\t".join(cell.text for cell in row.cells)
                    for row in block.rows
                )
    except (BadZipFile, OSError, PackageNotFoundError, ValueError) as exc:
        raise DocumentExtractionError(
            f"Could not extract DOCX '{filename}': {exc}"
        ) from exc

    page = ExtractedPage(
        page_number=None,
        text="\n".join(paragraphs_and_rows),
    )
    return ExtractedDocument(source_filename=filename, pages=(page,))


def _extract_txt(path: Path) -> ExtractedDocument:
    try:
        content = path.read_bytes()
    except OSError as exc:
        raise DocumentExtractionError(
            f"Could not read TXT '{path.name}': {exc}"
        ) from exc
    return _extract_txt_bytes(path.name, content)


def _extract_txt_bytes(filename: str, content: bytes) -> ExtractedDocument:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentExtractionError(
            f"Could not read TXT '{filename}' as UTF-8: {exc}"
        ) from exc

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    page = ExtractedPage(page_number=None, text=text)
    return ExtractedDocument(source_filename=filename, pages=(page,))
